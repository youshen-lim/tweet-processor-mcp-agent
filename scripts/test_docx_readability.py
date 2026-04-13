"""
Test script to verify .docx file readability and article parsing.
"""

import sys
from pathlib import Path

# Add src to path
sys.path.insert(0, str(Path(__file__).parent.parent / 'src'))


def test_docx_readability(docx_path: Path):
    """Test if .docx file can be read and parsed."""
    
    print("=" * 80)
    print("DOCX FILE READABILITY TEST")
    print("=" * 80)
    print()
    
    # Check file exists
    print(f"📂 Checking file: {docx_path}")
    if not docx_path.exists():
        print(f"❌ ERROR: File not found!")
        return False
    
    file_size = docx_path.stat().st_size / (1024 * 1024)  # MB
    print(f"✓ File exists ({file_size:.2f} MB)")
    print()
    
    # Test python-docx library
    print("📚 Testing python-docx library...")
    try:
        from docx import Document
        print("✓ python-docx library available")
    except ImportError:
        print("❌ ERROR: python-docx not installed")
        print("   Install with: pip install python-docx")
        return False
    print()
    
    # Read .docx file
    print("📖 Reading .docx file...")
    try:
        doc = Document(docx_path)
        print(f"✓ File opened successfully")
        print(f"   Paragraphs: {len(doc.paragraphs)}")
    except Exception as e:
        print(f"❌ ERROR: Failed to read .docx file")
        print(f"   Error: {e}")
        return False
    print()
    
    # Extract text
    print("📝 Extracting text content...")
    try:
        full_text = []
        for paragraph in doc.paragraphs:
            full_text.append(paragraph.text)
        
        text = '\n'.join(full_text)
        print(f"✓ Text extracted successfully")
        print(f"   Total characters: {len(text):,}")
        print(f"   Total lines: {len(full_text):,}")
    except Exception as e:
        print(f"❌ ERROR: Failed to extract text")
        print(f"   Error: {e}")
        return False
    print()
    
    # Parse articles using existing parser
    print("🔍 Parsing articles using DocumentParser...")
    try:
        from mcp_servers.google_drive_server import DocumentParser
        parser = DocumentParser(text)
        articles = parser.parse()
        
        print(f"✓ Articles parsed successfully")
        print(f"   Total articles found: {len(articles)}")
    except Exception as e:
        print(f"❌ ERROR: Failed to parse articles")
        print(f"   Error: {e}")
        import traceback
        traceback.print_exc()
        return False
    print()
    
    # Verify article count
    print("📊 Article Verification:")
    expected_count = 15
    if len(articles) == expected_count:
        print(f"✅ SUCCESS: Found all {expected_count} articles!")
    else:
        print(f"⚠️  WARNING: Expected {expected_count} articles, found {len(articles)}")
    print()
    
    # Display article summary
    print("📋 Article Summary:")
    print("-" * 80)
    
    total_words = 0
    for article in sorted(articles, key=lambda a: a.number):
        title_preview = article.title[:60] + "..." if len(article.title) > 60 else article.title
        url_preview = article.url[:50] + "..." if len(article.url) > 50 else article.url
        
        print(f"Article #{article.number:2d}:")
        print(f"  Title: {title_preview}")
        print(f"  URL:   {url_preview}")
        print(f"  Words: {article.word_count:,}")
        print()
        
        total_words += article.word_count
    
    print("-" * 80)
    print(f"Total word count: {total_words:,}")
    print()
    
    # Final verdict
    print("=" * 80)
    if len(articles) == expected_count:
        print("✅ READABILITY TEST PASSED")
        print("=" * 80)
        print()
        print("The .docx file is ready for migration!")
        print()
        print("Next step:")
        print(f"  python scripts/migrate_from_google_drive.py --input {docx_path}")
        print()
        return True
    else:
        print("⚠️  READABILITY TEST COMPLETED WITH WARNINGS")
        print("=" * 80)
        print()
        print(f"Expected {expected_count} articles but found {len(articles)}")
        print("Please review the article summary above.")
        print()
        return False


def main():
    import argparse
    
    parser = argparse.ArgumentParser(description="Test .docx file readability")
    parser.add_argument(
        '--input',
        type=str,
        default='data/articles.md.docx',
        help='Path to .docx file (default: data/articles.md.docx)'
    )
    
    args = parser.parse_args()
    
    # Get path
    repo_root = Path(__file__).parent.parent
    docx_path = repo_root / args.input
    
    # Run test
    success = test_docx_readability(docx_path)
    
    sys.exit(0 if success else 1)


if __name__ == "__main__":
    main()

